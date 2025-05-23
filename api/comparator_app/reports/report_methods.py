# import os
# import pandas as pd
# from openpyxl import Workbook, load_workbook
# from openpyxl.styles import PatternFill
# from openpyxl.utils import get_column_letter

# # Assuming these are defined in your comparator_app.reports.styles
# from comparator_app.reports.styles import red_fill, green_fill, bold_font 
# from comparator_app.utils.colors import Colors
# from comparator_app.configurator import decimal # Used for numerical comparison in reports

# # Assuming these are defined in your comparator_app.reports.report_utils
# from comparator_app.reports.report_utils import (
#     apply_border,
#     column_widths,
#     highlight_cells,
#     has_not_expected_or_missed,
#     has_not_expected,
#     is_positive_number,
#     is_zero
# )


# def generate_excel_report(records: list, output_file: str, highlight_rules: list = None, bold_header: bool = False):
#     """
#     Generates a generic Excel report from a list of dictionaries.

#     Args:
#         records (list): List of dictionaries, where each dictionary represents a row.
#         output_file (str): Full path and name of the output Excel file.
#         highlight_rules (list, optional): List of dictionaries defining conditional formatting rules.
#                                          Each rule dictates how cells should be highlighted.
#         bold_header (bool, optional): If True, the header row will be bold. Defaults to False.
#     """
#     if not records:
#         Colors.colored_print(f'No data to generate report for: {output_file}', 'WARNING')
#         return

#     try:
#         df = pd.DataFrame(records)
#         df.to_excel(output_file, index=False)

#         wb = load_workbook(output_file)
#         ws = wb.active

#         column_widths(ws)
#         apply_border(ws)
        
#         # Create a mapping from column name to its 1-based index for efficient lookup
#         col_name_to_idx = {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}

#         if highlight_rules:
#             for rule in highlight_rules:
#                 column = rule.get("column") # Highlight based on column name
#                 coords = rule.get("coords")   # Highlight based on (start_col, end_col) tuple
#                 condition = rule["condition"]
#                 fill_pass = rule.get("fill_pass")
#                 fill_fail = rule.get("fill_fail")

#                 if column:
#                     if column not in col_name_to_idx:
#                         Colors.colored_print(f"WARNING: Highlight rule column '{column}' not found in report headers.", 'WARNING')
#                         continue
#                     col_idx = col_name_to_idx[column]
#                     highlight_cells(ws, (col_idx, col_idx), condition, fill_pass, fill_fail)
#                 elif coords:
#                     highlight_cells(ws, coords, condition, fill_pass, fill_fail)

#         if bold_header:
#             for cell in ws[1]:
#                 cell.font = bold_font

#         wb.save(output_file)
#         Colors.colored_print(f'REPORT saved at: {output_file}', 'OKGREEN')

#     except Exception as e:
#         Colors.colored_print(f'ERROR creating REPORT: {e}', "FAIL")


# def highlighted_report(file_name: str, sheet_name: str, records: list, output_file: str):
#     """
#     Generates a highlighted Excel report for sheet comparison, indicating matches/differences.

#     Args:
#         file_name (str): Name of the file being reported.
#         sheet_name (str): Name of the sheet being reported.
#         records (list): List of dictionaries containing highlighted row data.
#         output_file (str): Full path and name of the output Excel file.
#     """
#     wb = Workbook()
#     ws = wb.active
#     ws.title = sheet_name

#     headers = ["Key", "Value File 1", "Value File 2", "Match", "Max_difference", "Difference"]
#     ws.append(headers)
    
#     pass_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid") # Green
#     fail_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid") # Red

#     for record in records:
#         key_raw = record.get('key', "")
        
#         # Handle tuple keys and 'missing' strings
#         if isinstance(key_raw, tuple):
#             key = " | ".join(str(item) for item in key_raw)
#         elif "missing" in str(key_raw).lower():
#             key = str(key_raw)
#         else:
#             key = str(key_raw) # Fallback for other types

#         value_file1 = str(record.get('value_file1', ''))
#         value_file2 = str(record.get('value_file2', ''))
#         match = record.get('match', "FAIL") # Default to FAIL if match status is missing
#         difference = str(record.get('difference', ''))
#         max_diff = record.get('max_diff', '')
        
#         row_data = [key, value_file1, value_file2, match, max_diff, difference]
#         ws.append(row_data)
        
#         current_row_idx = ws.max_row
        
#         # Apply fill based on 'Match' status
#         fill = pass_fill if match == "PASS" else fail_fill
#         for col in range(1, len(headers) + 1):
#             ws.cell(row=current_row_idx, column=col).fill = fill
            
#         # Override with red fill if there's a significant numerical difference
#         try:
#             if isinstance(max_diff, (int, float)) and abs(max_diff) > 10 ** -decimal:
#                 for col in range(1, len(headers) + 1):
#                     ws.cell(row=current_row_idx, column=col).fill = fail_fill
#         except Exception:
#             pass # Ignore if max_diff is not a comparable number

#     apply_border(ws) # Apply borders to all cells

#     # Auto-width columns
#     for col in ws.columns:
#         max_length = 0
#         column = get_column_letter(col[0].column)
#         for cell in col:
#             try: # Handle potential errors with cell.value
#                 if cell.value is not None:
#                     max_length = max(max_length, len(str(cell.value)))
#             except Exception:
#                 pass
#         ws.column_dimensions[column].width = max_length + 2 # Add some padding

#     wb.save(output_file)


# def total_report(records: list, output_file: str):
#     """
#     Generates the total summary Excel report with specific highlighting rules.

#     Args:
#         records (list): List of dictionaries containing the total comparison results.
#         output_file (str): Full path and name of the output Excel file.
#     """
#     generate_excel_report(
#         records=records,
#         output_file=output_file,
#         bold_header=True,
#         highlight_rules=[
#             {
#                 "column": "file_name", # Apply to file_name column
#                 "condition": has_not_expected_or_missed,
#                 "fill_pass": red_fill # Highlights red if file was missed or not expected
#             },
#             {
#                 "column": "sheet_name", # Apply to sheet_name column
#                 "condition": has_not_expected,
#                 "fill_pass": red_fill # Highlights red if sheet was not expected
#             },
#             {
#                 "column": "pass", # Apply to 'pass' column
#                 "condition": is_positive_number, # If 'pass' count is positive
#                 "fill_pass": green_fill,
#                 "fill_fail": red_fill
#             },
#             {
#                 "coords": (7, 10), # Apply to columns 7 through 10 (number_fail to max_difference)
#                 "condition": is_zero, # If values are zero
#                 "fill_pass": green_fill,
#                 "fill_fail": red_fill
#             }
#         ]
#     )

# The original `difference_report` functions were commented out.
# If you need a separate difference report, you can uncomment and adjust one of them.
# The `highlighted_report` already serves a similar purpose by showing differences.

import os
import pandas as pd
from openpyxl import load_workbook
from comparator_app.reports.report_utils import apply_border, column_widths, highlight_cells, has_not_expected_or_missed, has_not_expected, is_positive_number, is_zero
from comparator_app.reports.styles import red_fill, green_fill, bold_font
from comparator_app.utils.colors import Colors
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from comparator_app.configurator import decimal


def highlighted_report(file_name, sheet_name, records, output_file):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    headers = ["Key", "Value File 1", "Value File 2", "Match", "Max_difference", "Difference"]
    ws.append(headers)
    

    pass_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    fail_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    for record in records:
         
        key_raw = record.get('key', "")
    
        if isinstance(key_raw, str) or "missing" in str(key_raw).lower():
            key = str(key_raw)
        else:
            key = " | ".join(key_raw)
        value_file1 = str(record.get('value_file1', []))
        value_file2 = str(record.get('value_file2', []))
        match = record.get('match', "FAIL")
        difference = str(record.get('difference', []))
        max_diff = record.get('max_diff', "")
        
        row = [key, value_file1, value_file2, match, difference, max_diff]
        ws.append(row)
        
        fill = None
        try:
            if isinstance(max_diff, (int, float)) and abs(max_diff) > 10 ** -decimal:
                fill = fail_fill
        except (TypeError, ValueError, NameError) as e:
            print(f"Warning: max_diff check failed - {e}")

        if fill:
            for col in range(1, len(headers)+1):
                ws.cell(row=ws.max_row, column=col).fill = fill

        fill = pass_fill if match == "PASS" else fail_fill
        for col in range(1, len(headers)+1):
            ws.cell(row=ws.max_row, column=col).fill = fill
            
        apply_border(ws)

    # Auto-width
    for col in ws.columns:
        max_length = max(len(str(cell.value)) if cell.value else 0 for cell in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max_length + 2

    wb.save(output_file)



def total_report(records, output_file):
    generate_excel_report(
        records=records,
        output_file=output_file,
        bold_header=True,
        highlight_rules=[
            {
                "coords": (2, 2),
                "condition": has_not_expected_or_missed,
                "fill_pass": red_fill
            },
            {
                "coords": (1, 1),
                "condition": has_not_expected,
                "fill_pass": red_fill
            },
            {
                "coords": (6, 6),
                "condition": is_positive_number,
                "fill_pass": green_fill,
                "fill_fail": red_fill
            },
            {
                "coords": (7, 10),
                "condition": is_zero,
                "fill_pass": green_fill,
                "fill_fail": red_fill
            }
        ]
    )

# def difference_report(records, output_file, decimal):
#     generate_excel_report(
#         records=records,
#         output_file=output_file,
#         highlight_rules=[
#             {
#                 "coords": "max_diff",
#                 "condition": lambda v: isinstance(v, (int, float)) and abs(v) > 10 ** -decimal,
#                 "fill_pass": red_fill,
#                 "fill_fail": green_fill
#             }
#         ]
#     )

def generate_excel_report(records, output_file: str, highlight_rules: list = None, bold_header: bool = False):
    try:
        df = pd.DataFrame(records)
        df.to_excel(output_file, index=False)

        wb = load_workbook(output_file)
        ws = wb.active

        column_widths(ws)
        apply_border(ws)
        col_name_to_idx = {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}

        if highlight_rules:
            for rule in highlight_rules:
                column = rule.get("column")
                coords = rule.get("coords")
                condition = rule["condition"]
                fill_pass = rule.get("fill_pass")
                fill_fail = rule.get("fill_fail")

                if column:
                    if column not in col_name_to_idx:
                        continue
                    col_idx = col_name_to_idx[column]
                    highlight_cells(ws, (col_idx, col_idx), condition, fill_pass, fill_fail)
                elif coords:
                    highlight_cells(ws, coords, condition, fill_pass, fill_fail)

        if bold_header:
            for cell in ws[1]:
                cell.font = bold_font

        wb.save(output_file)
        Colors.colored_print(f'REPORT saved at: {output_file}', 'OKGREEN')

    except Exception as e:
        Colors.colored_print(f'ERROR creating REPORT: {e}', "FAIL")


# import pandas as pd
# from openpyxl import load_workbook
# from comparator_app.reports.report_utils import apply_border,column_widths,highlight_cells,has_not_expected_or_missed,has_not_expected,is_positive_number,is_zero
# from comparator_app.reports.styles import red_fill, green_fill, bold_font
# from comparator_app.utils.colors import Colors

# def highlighted_report(records, output_file):
#     generate_excel_report(
#         records=records,
#         output_file=output_file,
#         highlight_rules=[
#             {
#                 # "coords": (4, 4),
#                 "coords": 'match',
#                 "condition": lambda v: v == "PASS",
#                 "fill_pass": green_fill,
#                 "fill_fail": red_fill
#             }
#         ]
#     )

# def highlighted_report(records: list, sheet_name: str, output_folder: str):
#     # Prepare file path
#     file_name = f"MOCK_DATA_{sheet_name}.xlsx"
#     output_path = os.path.join(output_folder, file_name)

#     # Convert records into DataFrame
#     df = pd.DataFrame(records)

#     # Save DataFrame to Excel
#     df.to_excel(output_path, index=False, sheet_name=sheet_name)

#     # Load workbook and style
#     wb = load_workbook(output_path)
#     ws = wb.active

#     column_widths(ws)
#     apply_border(ws)

#     highlight_cells(
#         ws,
#         column_name="match",
#         condition=lambda v: v == "PASS",
#         fill_pass=green_fill,
#         fill_fail=red_fill
#     )

#     for cell in ws[1]:
#         cell.font = bold_font

#     wb.save(output_path)
#     return output_path

# def total_report(records, output_file):
#     generate_excel_report(
#         records=records,
#         output_file=output_file,
#         bold_header=True,
#         highlight_rules=[
#             {
#                 "coords": (2, 2),
#                 "condition": has_not_expected_or_missed,
#                 "fill_pass": red_fill
#             },
#             {
#                 "coords": (1, 1),
#                 "condition": has_not_expected,
#                 "fill_pass": red_fill
#             },
#             {
#                 "coords": (6, 6),
#                 "condition": is_positive_number,
#                 "fill_pass": green_fill,
#                 "fill_fail": red_fill
#             },
#             {
#                 "coords": (7, 10),
#                 "condition": is_zero,
#                 "fill_pass": green_fill,
#                 "fill_fail": red_fill
#             }
#         ]
#     )

# def difference_report(records, output_file, decimal):
#     generate_excel_report(
#         records=records,
#         output_file=output_file,
#         highlight_rules=[
#             {
#                 "coords": "max_diff",
#                 "condition": lambda v: isinstance(v, (int, float)) and abs(v) > 10 ** -decimal,
#                 "fill_pass": red_fill,
#                 "fill_fail": green_fill
#             }
#         ]
#     )

# def generate_excel_report(
#     records,
#     output_file: str,
#     highlight_rules: list = None,
#     bold_header: bool = False
# ):
#     try:
#         df = pd.DataFrame(records)
#         df.to_excel(output_file, index=False)

#         wb = load_workbook(output_file)
#         ws = wb.active

#         column_widths(ws)
#         apply_border(ws)
#         col_name_to_idx = {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}

#         if highlight_rules:
#             for rule in highlight_rules:
#                 column = rule.get("column")
#                 if column not in col_name_to_idx:
#                     continue  

#                 col_idx = col_name_to_idx[column]
#                 condition = rule["condition"]
#                 fill_pass = rule.get("fill_pass")
#                 fill_fail = rule.get("fill_fail")

#                 highlight_cells(ws, (col_idx, col_idx), condition, fill_pass, fill_fail)

#         if bold_header:
#             for cell in ws[1]:
#                 cell.font = bold_font

#         wb.save(output_file)
#         Colors.colored_print(f'REPORT saved at: {output_file}', 'OKGREEN')

#     except Exception as e:
#         Colors.colored_print(f'ERROR creating REPORT: {e}', "FAIL")

